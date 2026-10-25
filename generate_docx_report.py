import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text, title="CATATAN METODOLOGIS"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (thick navy)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(9.5)
    run_text.font.italic = True
    run_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, header_bg="1F4E79"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    # Data rows
    for r_idx, row in enumerate(table.rows[1:]):
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFC"
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for i, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[i])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                    
    # Add subtle borders to all cells
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="D3D3D3"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1F4E79"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E5E5E5"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def generate_report():
    doc = Document()
    
    # Page setup - Margins 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles setup
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)
    
    # ==========================
    # TITLE & HEADER
    # ==========================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("LAPORAN KESIMPULAN RISET & PENGEMBANGAN")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    sub_run = sub_p.add_run("From First Pressure to Defensive Response: Investigasi Koordinasi Pergerakan Bertahan Kolektif dan Kaitannya dengan Hasil Serangan")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x55, 0x66, 0x77)
    
    # Meta table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_widths = [2.2, 4.3]
    meta_headers = ["Parameter / Dimensi Proyek", "Keterangan / Rincian Implementasi"]
    for i, title in enumerate(meta_headers):
        meta_table.rows[0].cells[i].text = title
        
    meta_data = [
        ("Domain Riset", "Sports Analytics, Football Tracking & Event Data Science"),
        ("Dataset Validasi", "SkillCorner Open Data (Match ID: 1886347 — Auckland FC vs Newcastle Jets)"),
        ("Kerangka Kerja & Bahasa", "Python 3.14 (Pandas, NumPy, Matplotlib, Pillow, SciPy, Pytest)"),
        ("Status Pengujian & Pipeline", "100% Passed (26/26 Tests Lulus, 340 Episode Terproses Lengkap)")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[row_idx + 1]
        row.cells[0].text = k
        row.cells[1].text = v
    format_table(meta_table, meta_widths, header_bg="2B579A")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ==========================
    # 1. EXECUTIVE SUMMARY
    # ==========================
    h1 = doc.add_heading("1. Ringkasan Eksekutif", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Proyek riset 'From First Pressure to Defensive Response' bertujuan untuk memodelkan dan mengevaluasi "
        "bagaimana tindakan bertahan pertama (First Defender) direspons secara kolektif oleh rekan setimnya (Second Defender dan supporting defenders), "
        "serta bagaimana dinamika respons pergerakan ini memengaruhi hasil akhir serangan lawan (Attacking Outcome) dan nilai ancaman bertahan (Defensive Value / EPV)."
    )
    doc.add_paragraph(
        "Berbeda dengan metrik kepanduan konvensional yang berfokus pada pemeringkatan individu pemain, penelitian ini menempatkan "
        "koordinasi taktis defensif sebagai unit analisis utama. Seluruh alur kerja telah diimplementasikan secara terstruktur, "
        "reprodusibel, berbasis data pelacakan frekuensi tinggi (10 FPS) dan dynamic events resmi SkillCorner, serta diverifikasi melalui pengujian otomatis dan inspeksi visual animasi."
    )
    
    add_callout(
        doc,
        "Prinsip Fundamental Riset yang Dikunci: Setelah Episode Start (onset MDM First Defender), seluruh pergerakan pemain bertahan "
        "diamati tanpa filter keberhasilan atau ambang batas MDM. Pergerakan yang lambat, pasif, atau keliru arah tetap dicatat "
        "sebagai data empiris untuk menganalisis faktor kegagalan pertahanan secara objektif.",
        title="KEPUTUSAN METODOLOGIS BAKU"
    )

    # ==========================
    # 2. METODOLOGI & LOGIKA SISTEM
    # ==========================
    h2 = doc.add_heading("2. Metodologi dan Logika Sistem Bertahan", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Arsitektur pipeline dibangun berdasarkan aturan-aturan metodologis yang ketat dan transparan:"
    )
    
    p = doc.add_paragraph()
    p.add_run("a. Batasan Wilayah Sendiri (Halfway Line Constraint):\n").bold = True
    p.add_run(
        "Sesuai kebutuhan taktis, perhitungan dan pencarian episode defensif dimulai secara ketat ketika bola melintasi garis tengah (x = 0) "
        "dan masuk ke wilayah pertahanan sendiri (own half). Hal ini memfokuskan analisis pada fase pertahanan teritori aktif dan mengeliminasi penguasaan bola lawan yang berputar-putar di wilayah bertahannya sendiri."
    )
    
    p = doc.add_paragraph()
    p.add_run("b. Deteksi Meaningful Defensive Movement (MDM) Tiga Layer:\n").bold = True
    p.add_run("Kandidat First Defender dievaluasi melalui hierarki bukti terpadu:\n")
    p.add_run("• Layer 1 (Movement Evidence): ").bold = True
    p.add_run("Kecepatan puncak v ≥ 1.5 m/s dan perpindahan posisi netral Δd ≥ 0.6 m dalam jendela geser 0.5 detik (5 frame).\n")
    p.add_run("• Layer 2 (Defensive Context): ").bold = True
    p.add_run("Terlibat dalam event pressing/pressure SkillCorner, atau mendekati bola sebesar Δdist ≥ 0.8 m, atau aktif dalam duel 1v1 langsung (jarak ke pembawa bola ≤ 3.5 m).\n")
    p.add_run("• Layer 3 (Relevance): ").bold = True
    p.add_run("Jarak awal ke bola ≤ 25.0 m dan tidak bergerak menjauhi bola lebih dari 3.0 m.")
    
    p = doc.add_paragraph()
    p.add_run("c. Prinsip Second Defender (Pressure-Cover-Balance):\n").bold = True
    p.add_run(
        "Dalam teori taktik sepak bola, pemain yang memberikan cover harus berada dalam jangkauan jarak langsung dengan penekan pertama. "
        "Oleh karena itu, kandidat Second Defender dibatasi maksimal berjarak 18.0 m dari First Defender pada saat Episode Start. "
        "Pemain di sayap seberang (> 18.0 m) berfungsi menjaga keseimbangan struktur (team balance), bukan bertindak sebagai cover langsung. "
        "Di antara kandidat dalam radius 18 m, sistem memilih pemain dengan onset pergerakan tercepat dan yang berhasil menutup jarak terdekat (d_min)."
    )

    p = doc.add_paragraph()
    p.add_run("d. Evaluasi Hasil Akhir Serangan (Attacking Outcome) & Nilai Bertahan (Defensive Value):\n").bold = True
    p.add_run("Setiap episode diklasifikasikan ke dalam 3 hasil utama:\n")
    p.add_run("• PROGRESSION: ").bold = True
    p.add_run("Serangan maju secara signifikan (dorongan bola ke depan ≥ 8.0 m), menembus lini pertahanan, atau menghasilkan tembakan (shot).\n")
    p.add_run("• NO_PROGRESSION: ").bold = True
    p.add_run("Serangan tertahan, bola dipaksa dioper kembali ke belakang (recycled back), atau bergerak lateral tanpa kemajuan teritorial.\n")
    p.add_run("• TURNOVER: ").bold = True
    p.add_run("Tim bertahan sukses merebut penguasaan bola melalui tekel, intersep, sapuan terarah, atau kehilangan penguasaan oleh penyerang.\n")
    p.add_run("• Defensive Value (DV): ").bold = True
    p.add_run("Dihitung secara matematis melalui DV = EPV_against_start - EPV_against_end. Nilai positif (DV > 0) menandakan ancaman lawan berhasil diturunkan.")

    # ==========================
    # 3. KASUS EVALUASI VISUAL & PENYEMPURNAAN
    # ==========================
    h3 = doc.add_heading("3. Analisis Kasus Visual dan Penyempurnaan Algoritma", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Berdasarkan inspeksi visual terhadap animasi tactical GIF pada pertandingan 1886347, teridentifikasi anomali penentuan peran pemain "
        "pada implementasi awal yang kemudian berhasil disempurnakan secara komprehensif:"
    )
    
    # Table Comparison
    comp_table = doc.add_table(rows=3, cols=5)
    comp_widths = [1.1, 1.2, 1.4, 1.4, 1.4]
    comp_headers = ["Episode", "Peran", "Sebelum Revisi", "Setelah Revisi", "Rasional Taktis Sepak Bola"]
    for i, title in enumerate(comp_headers):
        comp_table.rows[0].cells[i].text = title
        
    comp_rows_data = [
        ("EP_000001", "1st Defender\n2nd Defender", "K. Grozos (Midfield)\nE. Adams (Deep CB)", "D. Ingham (RB)\nC. Timmins (CM)", 
         "Ingham menjaga rapat penyerang L. Gillion di garis tepi (jarak 1.36 m). Timmins melakukan covering sprint dari half-space mendekat ke 3.95 m."),
        ("EP_000002", "1st Defender\n2nd Defender", "A. Šušnjar (CB)\nM. Natta (CB)", "C. Timmins (CM)\nT. Aquilina (LB/WB)", 
         "Timmins langsung menekan F. Gallegos di lini tengah (jarak 4.44 m). Aquilina menutup jalur di sekitarnya (jarak 4.96 m).")
    ]
    for row_idx, data in enumerate(comp_rows_data):
        row = comp_table.rows[row_idx + 1]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].text = text
    format_table(comp_table, comp_widths, header_bg="1F4E79")
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph()
    p.add_run("Penyebab Akar Anomali dan Solusi yang Telah Diterapkan:\n").bold = True
    p.add_run(
        "1. Jockeying vs Sprinting Bias: Pemain bertahan yang berhadapan 1v1 dengan bola di area sayap melakukan gerakan jockeying "
        "(menjaga jarak dan mengatur keseimbangan tubuh) dengan kecepatan 1.5 - 2.5 m/s. Karena pemain sudah sangat dekat (1.3 m), ia tidak mungkin "
        "menutup jarak lagi sebesar 0.8 m tanpa melanggar lawan. Sementara itu, gelandang di lini tengah yang sedang berlari pemulihan (recovering sprint) "
        "mencapai kecepatan 6.7 m/s. Algoritma lama yang mengurutkan kandidat berdasarkan kecepatan tertinggi secara keliru memilih pelari pemulihan alih-alih pemain yang sedang berduel.\n"
        "2. Solusi Seleksi First Defender: Menambahkan konteks duel langsung (direct duel containment) dan mengurutkan kandidat pada onset terawal "
        "berdasarkan jarak terdekat ke bola (proximity to ball), sehingga bek yang menempel bola selalu terpilih sebagai First Defender.\n"
        "3. Solusi Seleksi Second Defender: Radius cover diperluas menjadi 18.0 m untuk menangkap pergerakan gelandang yang berlari menutup dari half-space "
        "menuju sayap, serta memprioritaskan pemain yang aktif mendekati First Defender (d_min terkecil)."
    )

    # ==========================
    # 4. HASIL PEMROSESAN PERTANDINGAN 1886347
    # ==========================
    h4 = doc.add_heading("4. Hasil Pemrosesan dan Statistik Pertandingan 1886347", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Eksekusi pipeline secara menyeluruh pada pertandingan Auckland FC vs Newcastle United Jets FC (ID: 1886347) menghasilkan metrik kuantitatif sebagai berikut:"
    )
    
    stat_table = doc.add_table(rows=8, cols=3)
    stat_widths = [2.2, 1.8, 2.5]
    stat_headers = ["Parameter / Metrik", "Nilai Terhitung", "Keterangan Taktis"]
    for i, title in enumerate(stat_headers):
        stat_table.rows[0].cells[i].text = title
        
    stat_rows_data = [
        ("Total Episode Bertahan", "340 Episode", "Dimulai saat bola masuk ke wilayah sendiri"),
        ("Total Pergerakan Defender", "3.618 Pergerakan", "Seluruh pergerakan pemain outfield tercatat"),
        ("Second Defender Terdeteksi", "339 Episode (99.7%)", "Hanya 1 episode tanpa cover (isolated press)"),
        ("Respon Simultan (≤ 0.20s)", "228 Episode (67.1%)", "Menunjukkan tingkat kesadaran kolektif tim"),
        ("Hasil: NO_PROGRESSION", "172 Episode (50.6%)", "Serangan lawan berhasil ditahan/didorong mundur"),
        ("Hasil: PROGRESSION", "122 Episode (35.9%)", "Serangan lawan menembus atau mencetak peluang"),
        ("Hasil: TURNOVER", "46 Episode (13.5%)", "Penguasaan bola berhasil direbut langsung")
    ]
    for row_idx, data in enumerate(stat_rows_data):
        row = stat_table.rows[row_idx + 1]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].text = text
    format_table(stat_table, stat_widths, header_bg="2B579A")
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================
    # 5. PENYEMPURNAAN VISUALISASI GIF
    # ==========================
    h5 = doc.add_heading("5. Penyempurnaan Visualisasi Animasi Tactical GIF", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Untuk memastikan transparansi dan kemudahan validasi visual oleh tim pelatih maupun analis olahraga, visualisasi animasi GIF telah ditingkatkan:"
    )
    p = doc.add_paragraph()
    p.add_run("• Kecepatan & Kelengkapan Frame (10 FPS):\n").bold = True
    p.add_run("GIF tidak lagi hanya memperlihatkan snapshot akhir, melainkan memutar seluruh urutan dari saat bola melintasi garis tengah, fase jockeying, hingga serangan tuntas pada kecepatan native 10 FPS (100 ms per frame).\n")
    p.add_run("• Legend & Identifikasi Tim yang Jelas:\n").bold = True
    p.add_run("Tim bertahan (Newcastle Jets) dirender dengan titik biru muda/cyan dan nomor punggung, sedangkan tim penyerang (Auckland FC) menggunakan titik oranye/merah. Posisi bola divisualisasikan dengan titik hitam tegas beraksen kuning.\n")
    p.add_run("• Sorotan Khusus Pemain Kunci:\n").bold = True
    p.add_run("First Defender ditandai dengan lingkaran merah menyala ('1st Def'), dan Second Defender ditandai dengan lingkaran oranye ('2nd Def').\n")
    p.add_run("• Banner Informasi Real-Time:\n").bold = True
    p.add_run("Menampilkan waktu pertandingan, nama kedua defender, selisih waktu respon (latency onset), jarak cover minimum, serta label outcome akhir.")

    # ==========================
    # 6. STATUS PENGUJIAN DAN ARTEFAK DOKUMENTASI
    # ==========================
    h6 = doc.add_heading("6. Verifikasi Kualitas dan Struktur Artefak", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)
    for r in h6.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Seluruh kode program telah melalui audit otomatis dan pengujian integrasi yang ketat:"
    )
    p = doc.add_paragraph()
    p.add_run("• Suite Pengujian Otomatis (Pytest): ").bold = True
    p.add_run("Sebanyak 26 unit & integration test yang mencakup kalkulasi jarak, kecepatan, orientasi vektor, batas garis tengah, rumus Defensive Value, penanganan data hilang, dan alur pipeline penuh berhasil lulus 100% (26 passed).\n")
    p.add_run("• Artefak Data CSV yang Dihasilkan:\n").bold = True
    p.add_run("1. data/processed/1886347/episodes.csv (340 baris episode)\n")
    p.add_run("2. data/processed/1886347/defender_movements.csv (3.618 baris pergerakan)\n")
    p.add_run("3. data/processed/1886347/episode_defender_features.csv (340 baris fitur koordinasi)\n")
    p.add_run("4. data/processed/1886347/episode_validation.csv (340 baris log validasi)\n")
    p.add_run("• Artefak Visual & Laporan:\n").bold = True
    p.add_run("1. outputs/1886347/gifs/ (Animasi taktis 10 FPS)\n")
    p.add_run("2. outputs/1886347/trajectories/ (Plot lintasan 2D)\n")
    p.add_run("3. outputs/1886347/run_manifest.json (Metadata eksekusi & parameter lengkap)\n")
    p.add_run("4. docs/methodology.md & docs/algorithm.md (Dokumentasi resmi metodologi dan algoritma)")

    # ==========================
    # 7. KESIMPULAN & LANGKAH LANJUTAN
    # ==========================
    h7 = doc.add_heading("7. Kesimpulan dan Rekomendasi Riset Selanjutnya", level=1)
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(6)
    for r in h7.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        
    doc.add_paragraph(
        "Sistem pipeline riset 'From First Pressure to Defensive Response' kini telah sepenuhnya matang, stabil, dan sesuai dengan prinsip taktis sepak bola modern. "
        "Algoritma ini terbukti mampu membedakan secara akurat antara duel jockeying di garis pertahanan depan dengan lari pemulihan dari lini kedua, "
        "serta mampu mengidentifikasi pemain cover secara kontekstual."
    )
    doc.add_paragraph(
        "Rekomendasi pengembangan riset selanjutnya meliputi:\n"
        "1. Penerapan Pipeline pada Multi-Match Dataset: Menjalankan pipeline pada seluruh pertandingan dalam dataset SkillCorner untuk membangun model klasifikasi koordinasi pertahanan berskala liga.\n"
        "2. Pemodelan Spatio-Temporal Lanjutan: Mengintegrasikan model Pitch Control (Voronoi/Potential Pitch Control) untuk mengukur volume ruang berbahaya yang berhasil ditutup oleh Second Defender.\n"
        "3. Analisis Profil Tim: Mengidentifikasi gaya bertahan tim (misalnya tim dengan respon cover sangat cepat vs tim yang mengandalkan garis pertahanan statis) dan korelasinya terhadap Expected Goals Conceded (xGA)."
    )
    
    # Save document
    out_dir = "reports"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Laporan_Kesimpulan_Defensive_Response.docx")
    doc.save(out_path)
    
    # Also save to root for easy user access
    root_path = "Laporan_Kesimpulan_Defensive_Response.docx"
    doc.save(root_path)
    
    print(f"Document successfully created at:\n1. {out_path}\n2. {root_path}")

if __name__ == "__main__":
    generate_report()
